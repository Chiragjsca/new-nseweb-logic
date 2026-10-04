import streamlit as st
import streamlit.components.v1 as components
import pandas as pd
import numpy as np
import re
import io
import os
import random
import zipfile
import gzip
import inspect
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink

# =====================================================================================
# 0. LOGIN GATE  (same two-mode gate as the original app)
# =====================================================================================
#   1) Streamlit secrets [auth] username/password  (preferred, keeps credentials off GitHub)
#   2) Hard-coded ADMIN_PASSWORD gate, used when no [auth] secrets are configured.
ADMIN_PASSWORD = "kano"

ADMIN_PASSWORD_ERROR_MESSAGES = [
    "Password इल्ले! 😅 इल्ले!, खम्मा घणी भाईसा, सॉरी। तुमसे सब कुछ हो पाएगा! यहां बहुत 🤪 दिमाग मत लगाओ, इस वेबसाइट को नहीं, 😂 इस गलत पासवर्ड को छोड़ दो!",
    "❌ Password इल्ले भाईसा! 😅 इल्ले! खम्मा घणी, सॉरी। तुम बाहुबली हो, तुमसे सब कुछ हो पाएगा! पर यहाँ फालतू 🤪 दिमाग मत लगाओ। अपनी सुंदर वेबसाइट को नहीं, 😂 इस सड़े हुए गलत पासवर्ड को छोड़ दो!",
    "❌ खम्मा घणी भाईसा, Password इल्ले! 😅 sorry! तुम तो मंगल ग्रह पर पानी खोज सकते हो, तुमसे सब कुछ हो पाएगा! पर यहाँ ज़्यादा 🤪 दिमाग मत लगाओ। इस सीधे-सादे वेबसाइट को नहीं, 😂 इस जाली पासवर्ड को छोड़ दो!",
    "❌ Password इल्ले! 😅 इल्ले! खम्मा घणी भाईसा, सॉरी। लोड मत लो, तुमसे सब कुछ हो पाएगा! पर यहाँ फालतू 🤪 दिमाग मत लगाओ। दुनिया छोड़ दो, मोक्ष पकड़ लो, पर पहले 😂 इस गलत पासवर्ड को छोड़ दो!",
    "❌ अरे भाईसा! Password इल्ले! 😅 खम्मा घणी, सॉरी। तुम चाहो तो सिस्टम हिला सकते हो, तुमसे सब कुछ हो पाएगा! पर यहाँ ज़्यादा 🤪 दिमाग मत लगाओ। इस निर्दोष वेबसाइट को नहीं, 😂 इस भूतिया गलत पासवर्ड को छोड़ दो!",
]

APP_TITLE = "Stocks Financial Data Merger & Formatter"


def _stretch():
    """Streamlit renamed use_container_width -> width='stretch'. Pick whichever
    this installed version understands so the app runs on old and new releases."""
    try:
        if "width" in inspect.signature(st.button).parameters:
            return {"width": "stretch"}
    except (TypeError, ValueError):
        pass
    return {"use_container_width": True}


def check_login():
    if st.session_state.get("authenticated", False):
        return True

    try:
        auth_cfg = st.secrets.get("auth", None)
    except Exception:
        auth_cfg = None

    # ---- Mode 1: secrets-based username/password ----
    if auth_cfg:
        st.title(f"🔒 {APP_TITLE} — Login")
        with st.form("login_form"):
            user = st.text_input("Username")
            pwd = st.text_input("Password", type="password")
            submitted = st.form_submit_button("Log in")
        if submitted:
            if user == auth_cfg.get("username") and pwd == auth_cfg.get("password"):
                st.session_state["authenticated"] = True
                st.rerun()
            else:
                st.error("Invalid username or password.")
        return False

    # ---- Mode 2: hard-coded ADMIN_PASSWORD gate ----
    st.markdown(
        "<p style='text-align: center; margin-top: 100px; color: Green; font-size: 18px;'>"
        f"📊 {APP_TITLE}</p>",
        unsafe_allow_html=True,
    )
    st.markdown(
        "<h1 style='text-align: center; margin-top: 0px; font-size: 20px;'>🔐 Admin Login</h1>",
        unsafe_allow_html=True,
    )
    _, col2, _ = st.columns([1, 1, 1])
    with col2:
        with st.form("admin_login_form"):
            pwd = st.text_input("Enter Password", type="password")
            submit = st.form_submit_button("Login", **_stretch())
            if submit:
                if pwd == ADMIN_PASSWORD:
                    st.session_state["authenticated"] = True
                    st.rerun()
                else:
                    st.error(random.choice(ADMIN_PASSWORD_ERROR_MESSAGES))
    st.markdown(
        f"<p style='text-align: center; color: gray; font-size: 14px; margin-top: 20px;'>"
        f"Data refreshed: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}</p>",
        unsafe_allow_html=True,
    )
    return False


# =====================================================================================
# 1. FILE CATALOGUE — which CSV is which category
#    stocks.csv -> index 0, "stocks (1).csv" -> index 1, ... "stocks (14).csv" -> 14
# =====================================================================================
TABS = [
    {"idx": 0,  "category": "Valuation Ratio",       "short": "Valuation"},
    {"idx": 1,  "category": "Profitability Ratios",  "short": "Profitability"},
    {"idx": 2,  "category": "Profitability Ratios",  "short": "Profitability"},
    {"idx": 3,  "category": "Profitability Ratios",  "short": "Profitability"},
    {"idx": 4,  "category": "Profitability Ratios",  "short": "Profitability"},
    {"idx": 5,  "category": "Growth Metrics",        "short": "Growth"},
    {"idx": 6,  "category": "Price Metrics",         "short": "Price"},
    {"idx": 7,  "category": "Price Metrics",         "short": "Price"},
    {"idx": 8,  "category": "Price Metrics",         "short": "Price"},
    {"idx": 9,  "category": "Price Metrics",         "short": "Price"},
    {"idx": 10, "category": "Technical Indicators",  "short": "Technical"},
    {"idx": 11, "category": "Technical Indicators",  "short": "Technical"},
    {"idx": 12, "category": "Profit & Loss Metrics", "short": "P&L"},
    {"idx": 13, "category": "Balance Sheet Metrics", "short": "Balance Sheet"},
    {"idx": 14, "category": "Balance Sheet Metrics", "short": "Balance Sheet"},
]
for _t in TABS:
    stem = "stocks" if _t["idx"] == 0 else f"stocks ({_t['idx']})"
    _t["stem"] = stem
    _t["file"] = f"{stem}.csv"
    # Excel sheet names: max 31 chars, none of []:*?/\
    _t["sheet"] = f"{_t['idx']} - {_t['short']}"[:31]
TAB_BY_IDX = {t["idx"]: t for t in TABS}

MASTER_SHEET_NAME = "Master_Dashboard"

# =====================================================================================
# 2. COLUMN DEFINITIONS
# =====================================================================================
# Columns present in every file (Symbol is generated from Company).
COMMON_COLUMNS = [
    "Symbol", "Company", "Sector", "Sub Sector", "Market Cap (in Cr)",
    "1D Return (%)", "1W Return (%)", "1M Return (%)", "Volume",
    "Price to Earning (P/E)",
]
TEXT_COLUMNS = {"Symbol", "Company", "Sector", "Sub Sector"}

# Category columns, in the exact order/spelling requested.
CATEGORY_COLUMNS = {
    "Valuation Ratio": ["Debt to Equity"],
    "Profitability Ratios": [
        "Earning Per Share (EPS)", "Dividend Per Share(DPS)", "Dividend Cover ratio",
        "Dividend Yield(%)", "Interest Coverage", "Face Value",
        "Promoter Holding (%)", "Public Holding",
    ],
    "Growth Metrics": [
        "Return Over % 1 month", "Return Over % Year to Date", "Return Over % 1 Year",
        "Return Over % 3 Years", "Return Over % 5 Years",
    ],
    "Price Metrics": [
        "Current Market Price", "VWAP", "Volume (in Lakhs)", "Open Price", "High Price",
        "Low Price", "Close Price", "Price Band Lower", "Price Band Higher",
        "52W Low", "52W High", "All time Low", "All time High",
        "Daily Volatility", "Annualized Volatility",
    ],
    "Technical Indicators": ["20 DMA", "50 DMA", "200 DMA"],
    "Profit & Loss Metrics": [
        "Total Income (in Lakhs)", "Total Expense (in Lakhs)", "Profit Before Tax (in Lakhs)",
        "Current Tax (in Lakhs)", "Deferred Tax (in Lakhs)", "Total Tax Expenses (in Lakhs)",
    ],
    "Balance Sheet Metrics": [
        "Total Equity (in Lakhs)", "Total Assets (in Lakhs)", "Current Assets (in Lakhs)",
        "Non-Current Assets (in Lakhs)", "Total Liabilities (in Lakhs)",
        "Current Liabilities (in Lakhs)", "Non-Current Liabilities (in Lakhs)",
        "Total Borrowings (in Lakhs)", "Total Revenue (in Lakhs)",
        "Long Term Borrowings (in Lakhs)",
    ],
}

CATEGORY_ORDER = list(CATEGORY_COLUMNS.keys())

# Raw CSV header -> requested column name. Matching ignores case and ALL whitespace,
# so "non- curr. Liabilities (in lakhs)" and "Non-Curr. Liabilities (In Lakhs)" both hit.
RAW_TO_TARGET_SPEC = {
    "Company": ["Company"],
    "Sector": ["Sector"],
    "Sub Sector": ["Sub Sector"],
    "Market Cap (in Cr)": ["Market Cap"],              # raw is in rupees -> converted to Cr
    "1D Return (%)": ["1D Return (%)"],
    "1W Return (%)": ["1W Return (%)"],
    "1M Return (%)": ["1M Return (%)"],
    "Volume": ["Volume"],
    "Price to Earning (P/E)": ["PE Ratio", "P/E", "PE"],
    "Debt to Equity": ["Debt/Equity", "Debt to Equity"],
    "Earning Per Share (EPS)": ["EPS"],
    "Dividend Per Share(DPS)": ["DPS"],
    "Dividend Cover ratio": ["Div. Cover", "Dividend Cover"],
    "Dividend Yield(%)": ["Div. Yield (%)", "Div. Yield", "Dividend Yield (%)"],
    "Interest Coverage": ["Interest Coverage", "Int. Coverage"],
    "Face Value": ["Face Value"],
    "Promoter Holding (%)": ["Promoter Holding (%)"],
    "Public Holding": ["Public Holding (%)", "Public Holding"],
    "Return Over % Year to Date": ["YTD Returns %", "YTD Return %"],
    "Return Over % 1 Year": ["1y Returns % (%)", "1Y Returns %", "1Y Return %"],
    "Return Over % 3 Years": ["3Y Return % (%)", "3Y Returns %", "3Y Return %"],
    "Return Over % 5 Years": ["5Y Returns % (%)", "5Y Return %", "5Y Returns %"],
    "VWAP": ["VWAP"],
    "Open Price": ["Open Price", "Open"],
    "High Price": ["High", "High Price"],
    "Low Price": ["D Low", "Low", "Low Price"],
    "Close Price": ["Close Price", "Close"],
    "Price Band Lower": ["Lower Price Band"],
    "Price Band Higher": ["Higher Price Band"],
    "52W Low": ["52WL", "52W Low"],
    "52W High": ["52WH", "52W High"],
    "All time Low": ["ATL"],
    "All time High": ["ATH"],
    "Daily Volatility": ["Daily Volatility"],
    "Annualized Volatility": ["Ann. Volatility", "Annualized Volatility"],
    "20 DMA": ["20 DMA"],
    "50 DMA": ["50 DMA"],
    "200 DMA": ["200 DMA"],
    "Total Income (in Lakhs)": ["Income (in lakhs)"],
    "Total Expense (in Lakhs)": ["Expense (in lakhs)"],
    "Current Tax (in Lakhs)": ["Curr. Tax (in lakhs)"],
    "Deferred Tax (in Lakhs)": ["Deff Tax (in lakhs)", "Def. Tax (in lakhs)"],
    "Total Tax Expenses (in Lakhs)": ["Total Tax (in lakhs)"],
    "Total Equity (in Lakhs)": ["Equity (in lakhs)"],
    "Total Assets (in Lakhs)": ["Total Assets (in lakhs)"],
    "Current Assets (in Lakhs)": ["Curr. Assets (in Lakhs)"],
    "Non-Current Assets (in Lakhs)": ["non-curr. assets (in lakhs)"],
    "Total Liabilities (in Lakhs)": ["Total Liabilities (in lakhs)"],
    "Current Liabilities (in Lakhs)": ["Curr. Liabilities (in lakhs)"],
    "Non-Current Liabilities (in Lakhs)": ["non- curr. Liabilities (in lakhs)"],
    "Total Borrowings (in Lakhs)": ["Total Debt (in lakhs)"],
}


def norm_key(text):
    return re.sub(r"\s+", "", str(text).strip().lower())


RAW_TO_TARGET = {
    norm_key(raw): target for target, raws in RAW_TO_TARGET_SPEC.items() for raw in raws
}
# The raw 'Market Cap' header is in rupees; only THIS header gets divided by 1e7.
RUPEES_PER_CRORE = 10_000_000
RAW_MARKET_CAP_KEY = norm_key("Market Cap")

# Columns that have no source in the CSVs. They are created empty so the sheet
# layout matches the requested list ("Long Term Borrowings" is empty by design).
# stocks (3) carries no extra column of its own, so it is the home for Interest Coverage.
BLANK_COLUMNS = {
    3: ["Interest Coverage"],
    14: ["Total Revenue (in Lakhs)", "Long Term Borrowings (in Lakhs)"],
}
# The files have no live-price column. When True, "Current Market Price" = Close Price.
CURRENT_PRICE_FROM_CLOSE = True

# ---- Number formats ----
F_AMT = "#,##0.00"
F_QTY = "#,##0"
F_PCT = '0.00"%"'
F_RATIO = "0.00"
COLUMN_FORMAT = {
    "Market Cap (in Cr)": F_AMT, "Volume": F_QTY,
    "1D Return (%)": F_PCT, "1W Return (%)": F_PCT, "1M Return (%)": F_PCT,
    "Price to Earning (P/E)": F_RATIO, "Debt to Equity": F_RATIO,
    "Dividend Cover ratio": F_RATIO, "Dividend Yield(%)": F_PCT, "Interest Coverage": F_RATIO,
    "Promoter Holding (%)": F_PCT, "Public Holding": F_PCT,
    "Return Over % 1 month": F_PCT, "Return Over % Year to Date": F_PCT,
    "Return Over % 1 Year": F_PCT, "Return Over % 3 Years": F_PCT, "Return Over % 5 Years": F_PCT,
    "Daily Volatility": F_RATIO, "Annualized Volatility": F_RATIO,
}
RETURN_COLUMNS = [
    "1D Return (%)", "1W Return (%)", "1M Return (%)", "Return Over % 1 month",
    "Return Over % Year to Date", "Return Over % 1 Year", "Return Over % 3 Years",
    "Return Over % 5 Years",
]

CATEGORY_FILL = {
    "Common": "D9D9D9", "Links": "D0E0E3", "Valuation Ratio": "DDEBF7", "Profitability Ratios": "E2EFDA",
    "Growth Metrics": "FFF2CC", "Price Metrics": "FCE4D6", "Technical Indicators": "EAD1DC",
    "Profit & Loss Metrics": "D9E1F2", "Balance Sheet Metrics": "E4DFEC",
}


def column_category(col):
    if col in HYPERLINK_SPECS:
        return "Links"
    if col in COMMON_COLUMNS:
        return "Common"
    for cat, cols in CATEGORY_COLUMNS.items():
        if col in cols:
            return cat
    return "Common"


def column_format(col, series=None):
    if col in HYPERLINK_SPECS:
        return "General"      # these cells hold formulas, so no Text format
    if col in TEXT_COLUMNS:
        return "@"
    if col in COLUMN_FORMAT:
        return COLUMN_FORMAT[col]
    return F_AMT


# Clickable link columns on the Master sheet, written as =HYPERLINK() formulas that read
# that row's Symbol cell. Same links as the original app.
# Format: "Column": (url_prefix, url_suffix, text_shown, append_symbol_to_text)
HYPERLINK_SPECS = {
    "NSE Chart": ("https://www.nseindia.com/get-quotes/equity?symbol=", "", "🟢", False),
    "NewW NSE": ("https://marketlens.nseindia.com/stocks/", "", "Nnse ", True),
    "Trading View": ("https://www.tradingview.com/symbols/", "", "Tre ", True),
    "History Data": ("https://www.equitypandit.com/historical-data/", "", "his ", True),
    "Chartlink": ("https://chartink.com/stocks/", ".html", "CL ", True),
    "Screener": ("https://www.screener.in/company/", "", "Scr ", True),
    "Marketsmith": ("https://marketsmithindia.com/mstool/eval/", "/evaluation.jsp", "ms ", True),
    "Zerodha": ("https://zerodha.com/markets/stocks/NSE/", "", "Z ", True),
}
LINK_COLUMNS = list(HYPERLINK_SPECS)

MASTER_ORDER = list(COMMON_COLUMNS)
for _cat in CATEGORY_ORDER:
    for _c in CATEGORY_COLUMNS[_cat]:
        if _c not in MASTER_ORDER:
            MASTER_ORDER.append(_c)
# NSE Chart dot sits right after Company (NewW NSE next to it); the other links go at the far right.
MASTER_ORDER.insert(MASTER_ORDER.index("Company") + 1, "NSE Chart")
MASTER_ORDER.insert(MASTER_ORDER.index("NSE Chart") + 1, "NewW NSE")   # right beside NSE Chart
MASTER_ORDER += [c for c in LINK_COLUMNS if c not in ("NSE Chart", "NewW NSE")]


# =====================================================================================
# 3. LOADING / STANDARDISING
# =====================================================================================
SYMBOL_RX = re.compile(r"\(([^()]*)\)\s*$")


def extract_symbol(company):
    """'Reliance Industries Limited (RELIANCE)' -> 'RELIANCE'.
    Uses the LAST bracket group so 'GAIL (India) Limited (GAIL)' -> 'GAIL'."""
    if company is None or (isinstance(company, float) and pd.isna(company)):
        return None
    m = SYMBOL_RX.search(str(company))
    if not m:
        return None
    sym = m.group(1).strip()
    return sym or None


def read_csv_bytes(file_bytes):
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return pd.read_csv(io.BytesIO(file_bytes), encoding=enc)
        except UnicodeDecodeError:
            continue
    return pd.read_csv(io.BytesIO(file_bytes), encoding="latin-1", encoding_errors="replace")


def clean_dataframe(df):
    """Drops fully-empty rows/columns of the RAW file (done before the requested
    empty placeholder columns are added, so those are never dropped)."""
    df = df.dropna(axis=0, how="all").dropna(axis=1, how="all")
    return df.reset_index(drop=True)


def standardise(df_raw, tab):
    """Raw CSV -> requested layout. Returns (dataframe, notes)."""
    notes = []
    df = df_raw.copy()
    df.columns = [str(c).strip() for c in df.columns]

    company_src = next((c for c in df.columns if norm_key(c) == "company"), None)
    if company_src is None:
        raise ValueError("no 'Company' column found")

    # Unit conversion keyed on the RAW header (so it can never be applied twice).
    for c in list(df.columns):
        if norm_key(c) == RAW_MARKET_CAP_KEY:
            df[c] = pd.to_numeric(df[c], errors="coerce") / RUPEES_PER_CRORE

    rename, seen = {}, set()
    for c in df.columns:
        target = RAW_TO_TARGET.get(norm_key(c))
        if target and target not in seen:
            rename[c] = target
            seen.add(target)
    df = df.rename(columns=rename)
    df = df.loc[:, ~df.columns.duplicated()]

    unmapped = [c for c in df.columns if c not in MASTER_ORDER]
    if unmapped:
        notes.append("Columns kept as-is (no rename rule): " + ", ".join(unmapped))

    # Numeric coercion for every non-text column
    for c in df.columns:
        if c in TEXT_COLUMNS or c in unmapped:
            continue
        if df[c].dtype == object:
            df[c] = pd.to_numeric(df[c].astype(str).str.replace(",", "").str.strip(), errors="coerce")

    # Symbol = text inside the last (...) of Company, placed LEFT of Company
    symbols = df["Company"].map(extract_symbol)
    missing = int(symbols.isna().sum())
    if missing:
        notes.append(f"{missing} Company value(s) have no (SYMBOL) in brackets — Symbol left blank.")
    df.insert(df.columns.get_loc("Company"), "Symbol", symbols)

    cat = tab["category"]
    if cat == "Price Metrics":
        if "Volume" in df.columns:
            df["Volume (in Lakhs)"] = df["Volume"] / 100_000
        if CURRENT_PRICE_FROM_CLOSE and "Close Price" in df.columns:
            df["Current Market Price"] = df["Close Price"]
    if cat == "Growth Metrics" and "1M Return (%)" in df.columns:
        df["Return Over % 1 month"] = df["1M Return (%)"]
    if tab["idx"] == 12 and "Total Income (in Lakhs)" in df.columns:
        # PBT = Income - Expense, only where an expense is actually reported (> 0)
        inc = df["Total Income (in Lakhs)"]
        exp = df.get("Total Expense (in Lakhs)", pd.Series(np.nan, index=df.index))
        df["Profit Before Tax (in Lakhs)"] = (inc - exp).where(exp > 0)
    for col in BLANK_COLUMNS.get(tab["idx"], []):
        if col not in df.columns:
            df[col] = np.nan

    wanted = COMMON_COLUMNS + CATEGORY_COLUMNS[cat]
    ordered = [c for c in wanted if c in df.columns]
    ordered += [c for c in df.columns if c not in ordered]
    return df[ordered], notes


@st.cache_data(show_spinner=False)
def load_standardised(file_bytes, tab_idx):
    raw = clean_dataframe(read_csv_bytes(file_bytes))
    return standardise(raw, TAB_BY_IDX[tab_idx])


def master_checks(mdf):
    """Cross-column sanity checks on the joined master."""
    out = []
    if {"20 DMA", "200 DMA"} <= set(mdf.columns):
        both = mdf[["20 DMA", "200 DMA"]].dropna()
        if len(both) and (both["20 DMA"] == both["200 DMA"]).mean() > 0.5:
            out.append(f"'20 DMA' equals '200 DMA' in {(both['20 DMA'] == both['200 DMA']).mean():.0%} of rows "
                       "— one of the two source files (stocks (10) / stocks (11)) looks wrong.")
    if {"Total Borrowings (in Lakhs)", "Total Assets (in Lakhs)"} <= set(mdf.columns):
        both = mdf[["Total Borrowings (in Lakhs)", "Total Assets (in Lakhs)"]].dropna()
        both = both[both["Total Assets (in Lakhs)"] > 0]
        if len(both) and (both.iloc[:, 0] > both.iloc[:, 1]).mean() > 0.2:
            out.append(f"'Total Borrowings' is larger than 'Total Assets' in "
                       f"{(both.iloc[:, 0] > both.iloc[:, 1]).mean():.0%} of rows — Total Debt in stocks (14) looks unreliable.")
    return out


def data_checks(df):
    """Flags numeric columns whose every value is identical (usually a broken export)."""
    out = []
    for c in df.columns:
        if c in TEXT_COLUMNS or not pd.api.types.is_numeric_dtype(df[c]):
            continue
        vals = df[c].dropna().unique()
        if len(df) > 1 and len(vals) == 0:
            out.append(f"'{c}' is empty for every row.")
        elif len(df) > 1 and len(vals) == 1:
            out.append(f"'{c}' has the same value ({vals[0]:g}) in every row.")
    return out


# =====================================================================================
# 4. MASTER DASHBOARD (all tabs joined on Symbol; first non-blank value wins)
# =====================================================================================
def build_master(processed):
    """processed: {tab_idx: DataFrame}. Returns (master_df with Symbol as a column, log)."""
    log, frames = [], []
    for t in TABS:
        d = processed.get(t["idx"])
        if d is None:
            continue
        if "Symbol" not in d.columns:
            log.append(f"{t['sheet']}: no Symbol column — skipped.")
            continue
        keep = d["Symbol"].notna() & (d["Symbol"].astype(str).str.strip() != "")
        if (~keep).any():
            log.append(f"{t['sheet']}: {int((~keep).sum())} row(s) without a Symbol were left out of the master.")
        d = d[keep].copy()
        d["Symbol"] = d["Symbol"].astype(str).str.strip()
        dup = int(d["Symbol"].duplicated().sum())
        if dup:
            log.append(f"{t['sheet']}: {dup} duplicate Symbol row(s) — first one kept.")
        frames.append(d.drop_duplicates("Symbol", keep="first").set_index("Symbol"))
    if not frames:
        return pd.DataFrame(columns=["Symbol"]), log
    master = frames[0]
    for f in frames[1:]:
        master = master.combine_first(f)
    master = master.sort_index()
    for lc in LINK_COLUMNS:
        master[lc] = ""      # filled with =HYPERLINK() formulas at export time
    cols = [c for c in MASTER_ORDER if c in master.columns and c != "Symbol"]
    cols += [c for c in master.columns if c not in cols]
    master = master[cols].reset_index()
    cols = [c for c in MASTER_ORDER if c in master.columns] + [c for c in master.columns if c not in MASTER_ORDER]
    return master[cols], log


# =====================================================================================
# 5. EXCEL EXPORT
# =====================================================================================
LINK_FONT = Font(color="0563C1", underline="single", bold=True)


def set_internal_link(cell, target):
    """target like "#'Main Tab'!A1" -> a real in-workbook link (location=, not a URL)."""
    cell.hyperlink = Hyperlink(ref=cell.coordinate, location=target.lstrip("#"), display=str(cell.value))
THIN = Side(style="thin", color="CCCCCC")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _autofit_widths(df):
    widths = []
    for c in df.columns:
        sample = df[c].head(500).dropna()
        if pd.api.types.is_numeric_dtype(df[c]):
            lens = [len(f"{v:,.2f}") for v in sample]
        else:
            lens = [len(str(v)) for v in sample]
        body = max(lens, default=0)
        widths.append(max(10, min(max(len(str(c)) + 3, body + 2), 45)))
    return widths


def write_table_sheet(ws, df, header_row, nav_target=None):
    """Header at header_row, data right under it, formats/filters/freeze applied."""
    n_cols = len(df.columns)
    if nav_target:
        link = ws.cell(row=1, column=1, value="⬆️ Main Tab")
        set_internal_link(link, nav_target)
        link.font = LINK_FONT

    for c, name in enumerate(df.columns, start=1):
        cell = ws.cell(row=header_row, column=c, value=str(name))
        cell.fill = PatternFill("solid", start_color=CATEGORY_FILL[column_category(name)],
                                end_color=CATEGORY_FILL[column_category(name)])
        cell.font = Font(name="Arial", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER

    body = df.astype(object).where(df.notna(), None)
    for row in body.itertuples(index=False, name=None):
        ws.append(list(row))

    first_data = header_row + 1
    last_row = header_row + len(df)

    link_cols = [(c, n) for c, n in enumerate(df.columns, start=1) if n in HYPERLINK_SPECS]
    if link_cols and "Symbol" in df.columns:
        sym_L = get_column_letter(list(df.columns).index("Symbol") + 1)
        link_font = Font(name="Arial", color="0563C1", underline="single")
        for c, name in link_cols:
            prefix, suffix, shown, with_symbol = HYPERLINK_SPECS[name]
            for r in range(first_data, last_row + 1):
                sym = f"{sym_L}{r}"
                label = f'"{shown}"&{sym}' if with_symbol else f'"{shown}"'
                cell = ws.cell(row=r, column=c, value=f'=HYPERLINK("{prefix}"&{sym}&"{suffix}",{label})')
                cell.font = link_font
                cell.alignment = Alignment(horizontal="center")

    for c, name in enumerate(df.columns, start=1):
        fmt = column_format(name)
        for r in range(first_data, last_row + 1):
            ws.cell(row=r, column=c).number_format = fmt

    for c, w in enumerate(_autofit_widths(df), start=1):
        ws.column_dimensions[get_column_letter(c)].width = w

    # Freeze header row + Symbol/Company columns
    lead = 0
    for name in df.columns:
        if name in ("Symbol", "Company", "NSE Chart"):
            lead += 1
        else:
            break
    freeze_col = max(lead, 1) + 1
    ws.freeze_panes = ws.cell(row=header_row + 1, column=freeze_col).coordinate
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(n_cols)}{max(last_row, header_row)}"

    # Green for positive / red for negative returns (font only, no fill)
    if last_row >= first_data:
        for c, name in enumerate(df.columns, start=1):
            if name in RETURN_COLUMNS:
                L = get_column_letter(c)
                rng = f"{L}{first_data}:{L}{last_row}"
                first = f"{L}{first_data}"
                ws.conditional_formatting.add(rng, FormulaRule(
                    formula=[f'AND(ISNUMBER({first}),{first}>0)'], font=Font(color="339966")))
                ws.conditional_formatting.add(rng, FormulaRule(
                    formula=[f'AND(ISNUMBER({first}),{first}<0)'], font=Font(color="FF0000")))


def build_workbook(tabs, master_df, master_order, include_tabs=True, source_names=None):
    """tabs: {tab_idx: df}. include_tabs=False -> only the Master sheet."""
    wb = Workbook()
    wb.remove(wb.active)
    exportable = [t for t in TABS if t["idx"] in tabs]

    if include_tabs:
        main = wb.create_sheet("Main Tab")
        main["A1"] = "📊 Stocks Financial Data — Main Tab"
        main["A1"].font = Font(bold=True, size=14)
        main["A3"] = "Click a name below to jump straight to that sheet."
        main["A3"].font = Font(italic=True)
        main["A4"] = "Header colours: " + " · ".join(CATEGORY_ORDER) + " (grey = common columns)."
        main["A4"].font = Font(italic=True)

        hdr = 5
        max_cols = max((len(tabs[t["idx"]].columns) for t in exportable), default=0)
        heads = ["Sheet", "Category", "Source file", "Rows", "Columns"] + [f"Column {i}" for i in range(1, max_cols + 1)]
        for c, h in enumerate(heads, start=1):
            main.cell(row=hdr, column=c, value=h).font = Font(bold=True)

        rows = [("⭐ " + MASTER_SHEET_NAME, "All categories joined on Symbol", "—", len(master_df),
                 len(master_order), list(master_order), f"#'{MASTER_SHEET_NAME}'!A1")]
        for t in exportable:
            d = tabs[t["idx"]]
            rows.append((f"➡️ {t['sheet']}", t["category"], (source_names or {}).get(t["idx"], "—"), len(d), len(d.columns),
                         list(d.columns), f"#'{t['sheet']}'!A1"))
        for i, (name, cat, src, n, nc, cols, target) in enumerate(rows, start=1):
            r = hdr + i
            cell = main.cell(row=r, column=1, value=name)
            set_internal_link(cell, target)
            cell.font = LINK_FONT
            main.cell(row=r, column=2, value=cat)
            main.cell(row=r, column=3, value=src)
            main.cell(row=r, column=4, value=n)
            main.cell(row=r, column=5, value=nc)
            for c, col in enumerate(cols, start=6):
                cc = main.cell(row=r, column=c, value=str(col))
                cc.fill = PatternFill("solid", start_color=CATEGORY_FILL[column_category(col)],
                                      end_color=CATEGORY_FILL[column_category(col)])
        main.column_dimensions["A"].width = 38
        main.column_dimensions["B"].width = 30
        main.column_dimensions["C"].width = 18
        main.column_dimensions["D"].width = 8
        main.column_dimensions["E"].width = 9
        for c in range(6, 6 + max_cols):
            main.column_dimensions[get_column_letter(c)].width = 24

    mdf = master_df[[c for c in master_order if c in master_df.columns]]
    ws = wb.create_sheet(MASTER_SHEET_NAME)
    write_table_sheet(ws, mdf, header_row=1)

    if include_tabs:
        for t in exportable:
            ws = wb.create_sheet(t["sheet"])
            write_table_sheet(ws, tabs[t["idx"]], header_row=2, nav_target="#'Main Tab'!A1")

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# =====================================================================================
# 6. PDF EXPORT (wide tables are split into column blocks; Symbol + Company repeat)
# =====================================================================================
PDF_MAX_ROWS_PER_TAB = 3000
PDF_COLS_PER_BLOCK = 9


def build_pdf_bytes(tabs_dict, title="Stocks Financial Data"):
    from reportlab.lib.pagesizes import landscape, A4
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak
    from reportlab.lib.styles import getSampleStyleSheet

    def fmt(v):
        if v is None or (isinstance(v, float) and np.isnan(v)):
            return ""
        if isinstance(v, (int, float, np.integer, np.floating)):
            return f"{v:,.2f}"
        return str(v)[:30]

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), title=title,
                            leftMargin=16, rightMargin=16, topMargin=16, bottomMargin=16)
    styles = getSampleStyleSheet()
    story = [Paragraph(title, styles["Title"]), Spacer(1, 10),
             Paragraph("Sheets in this file (click to jump):", styles["Heading2"])]
    for i, (name, df) in enumerate(tabs_dict.items()):
        story.append(Paragraph(f'<a href="#tab_{i}" color="blue">{name} — {len(df)} row(s)</a>', styles["Normal"]))
    story.append(PageBreak())

    avail = landscape(A4)[0] - 32
    for i, (name, df) in enumerate(tabs_dict.items()):
        story.append(Paragraph(f'<a name="tab_{i}"/>{name}', styles["Heading1"]))
        if df is None or df.empty:
            story.append(Paragraph("No rows after filtering.", styles["Normal"]))
            story.append(PageBreak())
            continue
        show = df.drop(columns=[c for c in LINK_COLUMNS if c in df.columns]).head(PDF_MAX_ROWS_PER_TAB)
        if len(df) > PDF_MAX_ROWS_PER_TAB:
            story.append(Paragraph(f"Showing first {PDF_MAX_ROWS_PER_TAB:,} of {len(df):,} rows.", styles["Italic"]))
        anchors = [c for c in ("Symbol", "Company") if c in show.columns]
        others = [c for c in show.columns if c not in anchors]
        blocks = [others[j:j + PDF_COLS_PER_BLOCK] for j in range(0, len(others), PDF_COLS_PER_BLOCK)] or [[]]
        for b, block in enumerate(blocks, start=1):
            cols = anchors + block
            if len(blocks) > 1:
                story.append(Paragraph(f"Columns block {b} of {len(blocks)}", styles["Italic"]))
            data = [cols] + [[fmt(v) for v in row] for row in show[cols].itertuples(index=False, name=None)]
            widths = []
            rest = avail - 60 * ("Symbol" in cols) - 130 * ("Company" in cols)
            for c in cols:
                widths.append(60 if c == "Symbol" else 130 if c == "Company" else max(rest / max(len(block), 1), 30))
            t = Table(data, repeatRows=1, colWidths=widths)
            t.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EAD1DC")),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 6),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]))
            story.append(t)
            story.append(PageBreak())
    doc.build(story)
    buf.seek(0)
    return buf


# =====================================================================================
# 7. INPUT HELPERS — zip/gzip ingestion, row selector, column sequencer
# =====================================================================================
class InMemoryFile:
    def __init__(self, name, data: bytes):
        self.name = name
        self._data = data

    def getvalue(self):
        return self._data


def extract_all_files(zip_bytes, _depth=0, _max_depth=3):
    """Walks a zip (and nested .zip / .gz members) and returns every leaf file."""
    results = []
    if _depth > _max_depth:
        return results
    try:
        with zipfile.ZipFile(io.BytesIO(zip_bytes)) as zf:
            for info in zf.infolist():
                if info.is_dir():
                    continue
                name = os.path.basename(info.filename)
                if not name or name.startswith("._") or "__MACOSX" in info.filename:
                    continue
                data = zf.read(info.filename)
                lower = name.lower()
                if lower.endswith(".zip"):
                    results.extend(extract_all_files(data, _depth + 1, _max_depth))
                elif lower.endswith(".gz"):
                    try:
                        results.append(InMemoryFile(name[:-3], gzip.decompress(data)))
                    except Exception:
                        pass
                else:
                    results.append(InMemoryFile(name, data))
    except zipfile.BadZipFile:
        pass
    return results


# A file is identified by the columns it contains, NOT by its name, so
# "stocks (12).csv" may hold any of the 15 data sets. These columns are unique to one slot:
SIGNATURE_TARGETS = {
    "Debt to Equity": 0,
    "Earning Per Share (EPS)": 1, "Dividend Per Share(DPS)": 1, "Dividend Cover ratio": 1,
    "Dividend Yield(%)": 1,
    "Face Value": 2,
    "Interest Coverage": 3,
    "Promoter Holding (%)": 4, "Public Holding": 4,
    "Return Over % Year to Date": 5, "Return Over % 1 Year": 5,
    "Return Over % 3 Years": 5, "Return Over % 5 Years": 5,
    "Open Price": 6, "High Price": 6, "Low Price": 6, "Close Price": 6,
    "Price Band Lower": 6, "Price Band Higher": 6,
    "VWAP": 7,
    "52W Low": 8, "52W High": 8, "All time Low": 8, "All time High": 8,
    "Daily Volatility": 9, "Annualized Volatility": 9,
    "200 DMA": 10,
    "20 DMA": 11, "50 DMA": 11,
    "Total Income (in Lakhs)": 12, "Total Expense (in Lakhs)": 12, "Current Tax (in Lakhs)": 12,
    "Deferred Tax (in Lakhs)": 12, "Total Tax Expenses (in Lakhs)": 12,
    "Total Equity (in Lakhs)": 13, "Total Assets (in Lakhs)": 13, "Current Assets (in Lakhs)": 13,
    "Non-Current Assets (in Lakhs)": 13, "Total Liabilities (in Lakhs)": 13,
    "Current Liabilities (in Lakhs)": 13,
    "Non-Current Liabilities (in Lakhs)": 14, "Total Borrowings (in Lakhs)": 14,
}
# A file with only the common columns (Company … P/E) has nothing of its own: it is slot 3,
# the Interest Coverage file, which has no source column.
BASE_ONLY_IDX = 3


@st.cache_data(show_spinner=False)
def detect_tab_idx(file_bytes):
    """Returns (slot_idx or None, reason). Reads the header row only."""
    try:
        head = read_csv_bytes_header(file_bytes)
    except Exception as e:
        return None, f"unreadable CSV ({e})"
    targets = {RAW_TO_TARGET.get(norm_key(c)) for c in head}
    if "Company" not in targets:
        return None, "no 'Company' column"
    slots = sorted({SIGNATURE_TARGETS[t] for t in targets if t in SIGNATURE_TARGETS})
    if not slots:
        return BASE_ONLY_IDX, "only common columns"
    if len(slots) > 1:
        return slots[0], f"columns from several data sets {slots}; treated as {slots[0]}"
    return slots[0], ""


def read_csv_bytes_header(file_bytes):
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return list(pd.read_csv(io.BytesIO(file_bytes), nrows=0, encoding=enc).columns)
        except UnicodeDecodeError:
            continue
    raise ValueError("cannot decode")


def parse_row_selector(text, n_rows):
    """'2,5,10-15' (1-indexed preview row numbers) -> set of 0-indexed positions."""
    indices = set()
    if not text:
        return indices
    for part in re.split(r"[,\s]+", text.strip()):
        if not part:
            continue
        try:
            if "-" in part:
                a, b = part.split("-", 1)
                a, b = int(a), int(b)
                if a > b:
                    a, b = b, a
                indices.update(i - 1 for i in range(a, b + 1) if 0 <= i - 1 < n_rows)
            else:
                i = int(part)
                if 0 <= i - 1 < n_rows:
                    indices.add(i - 1)
        except ValueError:
            continue
    return indices


def jump_to(anchor_id):
    st.session_state["scroll_target"] = anchor_id
    st.rerun()


def render_column_sequencer(state_key, current_columns, allow_delete=False, protected=None,
                            label="Column order", default_order=None):
    """Pick a column, then ◀ / ▶ (and optionally delete). Also accepts a typed,
    comma-separated sequence. Order persists in st.session_state[state_key]."""
    protected = protected or []
    if state_key not in st.session_state:
        st.session_state[state_key] = list(default_order) if default_order else list(current_columns)
    order = [c for c in st.session_state[state_key] if c in current_columns]
    for c in current_columns:
        if c not in order:
            order.append(c)
    st.session_state[state_key] = order
    if not order:
        return order

    st.caption(f"🔀 {label} — pick a column, then move it left/right" + (" or delete it:" if allow_delete else ":"))
    pick_col, left_col, right_col, del_col = st.columns([3, 1, 1, 1])
    with pick_col:
        pick = st.selectbox(label, options=order, key=f"{state_key}_pick", label_visibility="collapsed")
    with left_col:
        if st.button("◀ Left", key=f"{state_key}_left"):
            i = order.index(pick)
            if i > 0:
                order[i - 1], order[i] = order[i], order[i - 1]
                st.session_state[state_key] = order
                st.rerun()
    with right_col:
        if st.button("Right ▶", key=f"{state_key}_right"):
            i = order.index(pick)
            if i < len(order) - 1:
                order[i + 1], order[i] = order[i], order[i + 1]
                st.session_state[state_key] = order
                st.rerun()
    if allow_delete:
        with del_col:
            disabled = pick in protected
            if st.button("🗑 Delete", key=f"{state_key}_del", disabled=disabled,
                         help="This column is required and can't be deleted" if disabled else None):
                order.remove(pick)
                st.session_state[state_key] = order
                st.rerun()

    st.caption(" → ".join(order))

    type_col, apply_col = st.columns([5, 1])
    with type_col:
        typed = st.text_input("Or type the exact sequence here (comma-separated column names), then Apply:",
                              key=f"{state_key}_typed", placeholder=", ".join(order))
    with apply_col:
        st.write("")
        apply_clicked = st.button("Apply", key=f"{state_key}_apply")
    if apply_clicked:
        names = [t.strip() for t in typed.split(",") if t.strip()]
        if not names:
            st.warning("Type at least one column name before clicking Apply.")
        else:
            lookup = {c.strip().lower(): c for c in order}
            matched, unmatched = [], []
            for n in names:
                actual = lookup.get(n.lower())
                if actual and actual not in matched:
                    matched.append(actual)
                elif not actual:
                    unmatched.append(n)
            if not allow_delete:
                matched += [c for c in order if c not in matched]   # never silently drop data
            else:
                matched += [c for c in protected if c in order and c not in matched]
            if unmatched:
                st.warning(f"Not found (ignored): {', '.join(unmatched)}")
            if matched:
                st.session_state[state_key] = matched
                st.rerun()
    return order


# =====================================================================================
# 8. UI
# =====================================================================================
def main():
    st.set_page_config(page_title=APP_TITLE, layout="wide")
    if not check_login():
        st.stop()
    S = _stretch()

    st.title(f"📊 {APP_TITLE}")
    st.markdown('<div id="main_tab"></div>', unsafe_allow_html=True)
    st.caption(
        "Upload the 15 stocks CSV files in any order, with any names (or a ZIP of them). "
        "A **Symbol** column is generated from the text inside the brackets of **Company** and "
        "placed to its left."
    )

    # ---------------- uploads ----------------
    col_up1, col_up2 = st.columns(2)
    with col_up1:
        uploaded_files = st.file_uploader("Upload stocks CSV files", accept_multiple_files=True, type=["csv"])
    with col_up2:
        uploaded_zips = st.file_uploader("…or upload a ZIP containing the CSV files",
                                         accept_multiple_files=True, type=["zip"])
    candidates = list(uploaded_files) if uploaded_files else []
    if uploaded_zips:
        extracted_total = 0
        for z in uploaded_zips:
            ex = extract_all_files(z.getvalue())
            extracted_total += len(ex)
            candidates.extend(ex)
        st.caption(f"📦 Extracted {extracted_total} file(s) from {len(uploaded_zips)} zip archive(s).")

    st.markdown("---")

    # ---------------- checklist ----------------
    st.subheader("📋 File Checklist")
    st.caption("Files are recognised by the columns inside them, so the file names can be anything "
               "(e.g. 'stocks (12).csv' can hold any of the data sets).")
    by_tab, ignored = {}, []
    for f in candidates:
        if not f.name.lower().endswith(".csv"):
            ignored.append(f"{f.name} (not a .csv)")
            continue
        idx, why = detect_tab_idx(f.getvalue())
        if idx is None:
            ignored.append(f"{f.name} ({why})")
        else:
            by_tab.setdefault(idx, []).append(f)

    valid = {}
    status_cols = st.columns(3)
    for n, t in enumerate(TABS):
        col = status_cols[n % 3]
        files = by_tab.get(t["idx"])
        if files:
            chosen = files[0]
            if len(files) > 1:
                chosen = files[col.selectbox(
                    f"⚠️ {len(files)} files look like {t['category']} ({t['idx']}):", options=list(range(len(files))),
                    format_func=lambda i, _f=files: _f[i].name, key=f"select_{t['idx']}")]
            valid[t["idx"]] = chosen
            col.markdown(f"**✅ {t['idx']} · {t['category']}** <small style='color:green;'>"
                         f"← {chosen.name}</small>", unsafe_allow_html=True)
        else:
            col.markdown(f"**❌ {t['idx']} · {t['category']}** — "
                         f"<span style='color:#d9534f;font-weight:bold;'>Missing</span>", unsafe_allow_html=True)
    if ignored:
        st.warning("Ignored: " + "; ".join(sorted(ignored)))

    st.markdown("---")

    # ---------------- main-tab navigation ----------------
    st.subheader("🏠 Main Tab — Quick Navigation")
    st.caption("Click a section to jump to it. Each section has a '⬆️ Back to Main Tab' button.")
    nav_cols = st.columns(3)
    for n, t in enumerate(TABS):
        with nav_cols[n % 3]:
            if t["idx"] in valid:
                if st.button(f"➡️ {t['sheet']}", key=f"jump_{t['idx']}", **S):
                    jump_to(f"tab_{t['idx']}")
            else:
                st.button(f"🚫 {t['sheet']}", key=f"jump_disabled_{t['idx']}", disabled=True, **S,
                          help="Not available — upload the matching file first.")
    st.markdown("---")

    if not valid:
        if candidates:
            st.warning("⚠️ None of the uploaded files could be recognised as one of the 15 data sets.")
        _scroll()
        return

    # ---------------- per-tab tuning ----------------
    st.subheader("🛠️ Component Tuning & Data Previews")
    processed, all_notes = {}, {}
    for t in TABS:
        idx = t["idx"]
        if idx not in valid:
            continue
        f = valid[idx]
        try:
            df, notes = load_standardised(f.getvalue(), idx)
        except Exception as e:
            st.error(f"Could not read {f.name}: {e}")
            continue

        st.markdown(f'<div id="tab_{idx}"></div>', unsafe_allow_html=True)
        expanded = st.session_state.get("scroll_target") == f"tab_{idx}"
        with st.expander(f"{t['sheet']}  —  {t['category']}  ({f.name})", expanded=expanded):
            if st.button("⬆️ Back to Main Tab", key=f"back_{idx}"):
                jump_to("main_tab")

            cols = df.columns.tolist()
            ctrl_a, ctrl_b = st.columns(2)
            with ctrl_a:
                removals = st.multiselect(f"Columns to remove from {t['sheet']}:",
                                          options=[c for c in cols if c not in ("Symbol", "Company")],
                                          key=f"remove_{idx}")
            with ctrl_b:
                row_text = st.text_input("Rows to remove (row #s shown in the preview, e.g. 2,5,10-15):",
                                         key=f"rowsel_{idx}")
            d1, d2, d3 = st.columns([1, 1, 2])
            with d1:
                dedupe = st.checkbox("Remove duplicate rows", key=f"dedupe_{idx}")
            with d2:
                filt_col = st.selectbox("Remove rows matching a value in:", ["(none)"] + cols, key=f"filtercol_{idx}")
            exclude_vals = []
            if filt_col != "(none)":
                uniq = sorted(df[filt_col].dropna().astype(str).unique().tolist())[:500]
                with d3:
                    exclude_vals = st.multiselect(f"Value(s) to remove from '{filt_col}':", options=uniq,
                                                  key=f"filtervals_{idx}")

            out = df.drop(columns=removals)
            if dedupe:
                out = out.drop_duplicates()
            if exclude_vals and filt_col in out.columns:
                out = out[~out[filt_col].astype(str).isin(exclude_vals)]
            out = out.reset_index(drop=True)
            drop_rows = parse_row_selector(row_text, len(out))
            if drop_rows:
                out = out.drop(index=sorted(drop_rows)).reset_index(drop=True)

            order = render_column_sequencer(f"colorder_{idx}", out.columns.tolist(),
                                            label=f"Column order for {t['sheet']}",
                                            default_order=out.columns.tolist())
            out = out[order]
            processed[idx] = out
            all_notes[idx] = notes + data_checks(out)

            st.caption(f"Rows: {len(df)} original → {len(out)} after cleanup")
            for n in all_notes[idx]:
                st.caption(f"⚠️ {n}")
            st.dataframe(out.head(10), **S)

    if not processed:
        _scroll()
        return

    # ---------------- data notes ----------------
    master_df, master_log = build_master(processed)
    st.markdown("---")
    with st.expander("ℹ️ Data notes & checks (read before relying on the numbers)", expanded=False):
        st.markdown(
            "- **Symbol** is the text inside the last `( )` of *Company*.\n"
            "- **Market Cap (in Cr)** = the file's *Market Cap* ÷ 10,000,000 (the raw values are in rupees).\n"
            "- **Volume (in Lakhs)** = *Volume* ÷ 100,000 (added in the Price tabs).\n"
            + ("- **Current Market Price** has no source column in the files, so it is set equal to **Close Price**.\n"
               if CURRENT_PRICE_FROM_CLOSE else "- **Current Market Price** is left empty (no source column).\n")
            + "- **Empty by design / no source in the files:** Interest Coverage (stocks (3)), "
              "Profit Before Tax (only filled where an expense > 0 is reported), Total Revenue and "
              "**Long Term Borrowings** (stocks (14), always empty).\n"
            "- **Return Over % 1 month** is a copy of *1M Return (%)* (Growth tab).\n"
            "- A file can hold fewer companies than the others; the master keeps every Symbol found "
            "in any file and leaves the missing cells blank."
        )
        counts = {TAB_BY_IDX[i]["sheet"]: len(d) for i, d in processed.items()}
        st.write("Rows per sheet:", counts)
        for n in master_checks(master_df):
            st.warning(n)
        flagged = {TAB_BY_IDX[i]["sheet"]: n for i, n in all_notes.items() if n}
        if flagged:
            st.write("Checks that fired:")
            for sheet, ns in flagged.items():
                for n in ns:
                    st.write(f"- **{sheet}** — {n}")

    # ---------------- master columns ----------------
    st.markdown("---")
    st.subheader(f"🔀 {MASTER_SHEET_NAME} — column order & inclusion")
    st.caption("The master joins every sheet above on Symbol (first non-blank value wins). "
               "Reorder or delete columns here. 'Symbol' and 'Company' can't be deleted.")
    master_order = render_column_sequencer(
        "master_col_order", master_df.columns.tolist(), allow_delete=True,
        protected=["Symbol", "Company"], label=f"{MASTER_SHEET_NAME} columns",
        default_order=master_df.columns.tolist())
    st.markdown("---")

    # ---------------- execute ----------------
    if st.button("🚀 Execute Structural Consolidation", type="primary"):
        with st.spinner("Building Excel files…"):
            st.session_state["consolidation_result"] = {
                "output_bytes": build_workbook(processed, master_df, master_order, include_tabs=True,
                                               source_names={i: f.name for i, f in valid.items()}),
                "master_only_bytes": build_workbook(processed, master_df, master_order, include_tabs=False),
                "master_df": master_df,
                "master_order": list(master_order),
                "master_log": master_log,
                "processed": processed,
            }
        st.session_state.pop("all_tabs_pdf_bytes", None)
        st.session_state.pop("master_pdf_bytes", None)

    result = st.session_state.get("consolidation_result")
    if result:
        st.success("✅ Consolidation and Formatting Complete!")
        mdf, morder = result["master_df"], result["master_order"]
        if result["master_log"]:
            with st.expander(f"⚠️ {MASTER_SHEET_NAME}: {len(result['master_log'])} warning(s)"):
                for line in result["master_log"]:
                    st.write("- " + line)
        m1, m2, m3 = st.columns(3)
        m1.metric(f"{MASTER_SHEET_NAME} Symbols", len(mdf))
        m2.metric("Columns", len(morder))
        m3.metric("Duplicate Symbols", int(mdf["Symbol"].duplicated().sum()) if "Symbol" in mdf else 0)
        st.dataframe(mdf[[c for c in morder if c in mdf.columns]].head(20), **S)

        ts = datetime.now().strftime("%Y%m%d_%H%M")
        dl1, dl2 = st.columns(2)
        with dl1:
            st.download_button("📥 Download Formatted Master File", data=result["output_bytes"],
                               file_name=f"Master_Financial_Data_{ts}.xlsx",
                               mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                               type="primary", key="download_xlsx_btn")
        with dl2:
            st.download_button(f"📥 Download {MASTER_SHEET_NAME} Only (separate Excel)",
                               data=result["master_only_bytes"],
                               file_name=f"{MASTER_SHEET_NAME}_{ts}.xlsx",
                               mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                               key="download_master_only_xlsx_btn")

        st.markdown("---")
        st.subheader("📄 PDF Export")
        st.caption("A PDF can't hold live Excel filter dropdowns; it mirrors the rows/columns you've already "
                   "filtered above. Wide tables are split into column blocks (Symbol + Company repeat). "
                   "Use the .xlsx for live filtering.")
        p1, p2 = st.columns(2)
        with p1:
            if st.button("📄 Build All-Sheets PDF", key="build_all_pdf_btn"):
                try:
                    with st.spinner("Building All-Sheets PDF…"):
                        named = {TAB_BY_IDX[i]["sheet"]: d for i, d in result["processed"].items()}
                        st.session_state["all_tabs_pdf_bytes"] = build_pdf_bytes(
                            named, "Stocks Financial Data — All Sheets").getvalue()
                except ImportError:
                    st.error("PDF export needs the `reportlab` package (pip install reportlab).")
            if st.session_state.get("all_tabs_pdf_bytes"):
                st.download_button("📥 Download All-Sheets PDF", data=st.session_state["all_tabs_pdf_bytes"],
                                   file_name=f"All_Sheets_{ts}.pdf", mime="application/pdf",
                                   key="download_all_pdf_btn")
        with p2:
            if st.button(f"📄 Build {MASTER_SHEET_NAME} PDF", key="build_master_pdf_btn"):
                try:
                    with st.spinner("Building master PDF…"):
                        st.session_state["master_pdf_bytes"] = build_pdf_bytes(
                            {MASTER_SHEET_NAME: mdf[[c for c in morder if c in mdf.columns]]},
                            MASTER_SHEET_NAME).getvalue()
                except ImportError:
                    st.error("PDF export needs the `reportlab` package (pip install reportlab).")
            if st.session_state.get("master_pdf_bytes"):
                st.download_button(f"📥 Download {MASTER_SHEET_NAME} PDF", data=st.session_state["master_pdf_bytes"],
                                   file_name=f"{MASTER_SHEET_NAME}_{ts}.pdf", mime="application/pdf",
                                   key="download_master_pdf_btn")
    _scroll()


def _scroll():
    """Runs last so every anchor exists; consumes the pending scroll target once."""
    target = st.session_state.get("scroll_target")
    if target:
        st.session_state["scroll_target"] = None
        components.html(
            f"""<script>
            setTimeout(function() {{
                var el = window.parent.document.getElementById("{target}");
                if (el) {{ el.scrollIntoView({{behavior: "smooth", block: "start"}}); }}
            }}, 150);
            </script>""",
            height=0,
        )


if __name__ == "__main__":
    main()
